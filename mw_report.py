#!/usr/bin/env python3
"""
Builds the daily "Please verify MWs" message from the MW index Word doc and
every SharePoint workbook it links to. Files are read straight into memory
via the SharePoint REST API; nothing is saved to disk.

Auth: the first run opens Chrome so you can log in with Cisco SSO; the
session cookies are then stored in session.json and reused headlessly.
Chrome only reopens when that session has expired.

Setup (once):
  python3 -m pip install --user --break-system-packages -r requirements.txt
  python3 -m playwright install chromium

Usage:
  python3 mw_report.py                     # today's MWs
  python3 mw_report.py --date 2026-10-02   # a specific date
  python3 mw_report.py --inspect           # dump each workbook's structure
"""
import argparse
import datetime as dt
import difflib
import io
import json
import os
import re
import subprocess
import sys
import time
import warnings
import zipfile
from collections import Counter
from pathlib import Path
from urllib.parse import parse_qs, urlsplit
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

warnings.filterwarnings("ignore", module="openpyxl")

PROJECT_DIR = Path(__file__).parent
SESSION_FILE = PROJECT_DIR / "session.json"
CONFIG_FILE = PROJECT_DIR / "config.json"
MESSAGES_DIR = PROJECT_DIR / "messages"
LOGIN_TIMEOUT_S = 300
FETCH_ATTEMPTS = 3

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
GUID_RE = r"[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}"

DATE_COLS = ("es_mw_date", "mw_date")
TIME_COLS = ("es_mw_time", "v1_mw_time", "mw_time")
STATUS_COLS = ("rm_status", "rm status", "status")
# Column on the RT sheet that says whether the stack also gets KV migration.
KV_COLS = ("op_exec_md_mongo_to_kvservice",)
# "Oct 10, 2026 04:00-06:00 UTC" (some sheets drop the comma) -- the column the team schedules from, so it
# wins over the date/time columns whenever a sheet has it.
FORMAL_COLS = ("formal_date_time",)
FORMAL_RE = re.compile(r"([A-Z][a-z]{2}) (\d{2}),? (\d{4}) (\d{2}:\d{2}.*)")
# Every MW counts unless it was called off or scheduled outside the batch;
# statuses move on to Completed, Failed, etc. once the MW runs. Rows with no
# status at all were never put in the batch (mws-validation is 0 on them).
EXCLUDED_STATUSES = ("cancelled", "canceled", "scheduled manually")
SEPARATOR = "-" * 37


def log(msg: str):
    print(msg, file=sys.stderr)


class AuthRequired(Exception):
    pass


def first_line(e: Exception) -> str:
    # Playwright errors append a call log that includes the session cookies;
    # keep only the summary line so they never reach the message.
    return (str(e).strip().splitlines() or [type(e).__name__])[0]


# ---------------------------------------------------------------- SharePoint

def file_api_url(base: str, site: str, guid: str) -> str:
    return f"{base}{site}/_api/web/GetFileById('{guid}')/$value"


def direct_api_url(link: str) -> str | None:
    parts = urlsplit(link)
    sourcedoc = parse_qs(parts.query).get("sourcedoc", [None])[0]
    if not sourcedoc:
        return None
    path = re.sub(r"^/:[a-zA-Z]:/r/", "/", parts.path)
    idx = path.lower().find("/_layouts/")
    site = path[:idx] if idx != -1 else ""
    return file_api_url(f"{parts.scheme}://{parts.netloc}", site, sourcedoc.strip("{}"))


def is_login_page(resp) -> bool:
    return resp.status == 401 or "login.microsoftonline" in resp.url


def interactive_login(p, link: str):
    host = urlsplit(link).netloc
    browser = p.chromium.launch(headless=False)
    try:
        context = browser.new_context(
            storage_state=str(SESSION_FILE) if SESSION_FILE.exists() else None
        )
        page = context.new_page()
        page.goto(link, wait_until="domcontentloaded")
        log(f"Log in to Cisco SSO in the Chrome window (waiting for {host}) ...")
        deadline = time.time() + LOGIN_TIMEOUT_S
        while urlsplit(page.url).netloc != host:
            if time.time() > deadline:
                sys.exit("Login timed out.")
            page.wait_for_timeout(2000)
        page.wait_for_load_state("domcontentloaded")
        context.storage_state(path=str(SESSION_FILE))
        os.chmod(SESSION_FILE, 0o600)
        log("Login captured; session saved.")
    finally:
        browser.close()


class SharePointReader:
    def __init__(self, p):
        self.p = p
        self.req = None

    def _open(self):
        if self.req is None:
            self.req = self.p.request.new_context(storage_state=str(SESSION_FILE))

    def _resolve_sharing_link(self, link: str) -> str:
        # Short sharing links (/:x:/s/<Site>/<token>) carry no GUID; the
        # viewer page they open embeds it, so read it from there.
        resp = self.req.get(link, timeout=180000)
        if is_login_page(resp):
            raise AuthRequired(f"HTTP {resp.status} at {resp.url}")
        m = re.search(rf"(?i)(?:sourcedoc=|UniqueId[\"':=\s]+)(?:%7B|[{{\"'])*({GUID_RE})", resp.text())
        site = re.match(r"^/:[a-zA-Z]:/s/([^/]+)/", urlsplit(link).path)
        if not m or not site:
            raise RuntimeError(f"Could not resolve file id for sharing link {link}")
        parts = urlsplit(link)
        return file_api_url(f"{parts.scheme}://{parts.netloc}", f"/sites/{site.group(1)}", m.group(1))

    def _fetch(self, link: str) -> bytes:
        # Large workbooks sometimes drop mid-download ("aborted"); retry those.
        for attempt in range(1, FETCH_ATTEMPTS + 1):
            try:
                return self._fetch_once(link)
            except (AuthRequired, RuntimeError):
                raise
            except Exception as e:
                if attempt == FETCH_ATTEMPTS:
                    raise
                log(f"  retrying after: {first_line(e)}")
                time.sleep(2 * attempt)

    def _fetch_once(self, link: str) -> bytes:
        url = direct_api_url(link) or self._resolve_sharing_link(link)
        resp = self.req.get(url, timeout=180000)
        ctype = resp.headers.get("content-type", "")
        # An expired session on the REST API comes back as 403, not 401.
        if is_login_page(resp) or resp.status == 403 or "text/html" in ctype:
            raise AuthRequired(f"HTTP {resp.status} at {resp.url}")
        if not resp.ok:
            raise RuntimeError(f"HTTP {resp.status} reading {link}: {resp.text()[:300]}")
        return resp.body()

    def read(self, link: str) -> bytes:
        if not SESSION_FILE.exists():
            interactive_login(self.p, link)
        self._open()
        try:
            return self._fetch(link)
        except AuthRequired:
            self.close()
            interactive_login(self.p, link)
            self._open()
            try:
                return self._fetch(link)
            except AuthRequired as e:
                raise RuntimeError(
                    f"Still unauthorized after fresh login ({e}); check that your "
                    f"account has access to {link}"
                ) from None

    def close(self):
        if self.req is not None:
            self.req.dispose()
            self.req = None


# ------------------------------------------------------------------- Parsing

def docx_hyperlinks(data: bytes):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        rels_xml = z.read("word/_rels/document.xml.rels")
        doc_xml = z.read("word/document.xml")
    rels = {r.attrib["Id"]: r.attrib.get("Target", "") for r in ET.fromstring(rels_xml)}
    root = ET.fromstring(doc_xml)

    links = []
    for h in root.iter(f"{W_NS}hyperlink"):
        url = rels.get(h.attrib.get(f"{R_NS}id"), "")
        text = "".join(t.text or "" for t in h.iter(f"{W_NS}t")).strip()
        if url:
            links.append((text, url))
    # Field-code hyperlinks: <w:instrText> HYPERLINK "url" </w:instrText>
    instr = "".join(t.text or "" for t in root.iter(f"{W_NS}instrText"))
    for url in re.findall(r'HYPERLINK\s+"([^"]+)"', instr):
        links.append(("", url))

    seen, out = set(), []
    for text, url in links:
        if "sharepoint.com" in url and url not in seen:
            seen.add(url)
            out.append((text, url))
    return out


def norm(s: str) -> str:
    return " ".join(s.replace("\xa0", " ").split()).lower()


def header(ws) -> list[str]:
    row = next(ws.iter_rows(max_row=1, values_only=True), ())
    return [str(c).strip().lower() if c is not None else "" for c in row]


def first_col(hdr: list[str], names) -> int | None:
    for n in names:
        if n in hdr:
            return hdr.index(n)
    return None


def stack_tab(wb, preferred: str | None):
    if preferred and preferred in wb.sheetnames:
        return wb[preferred]
    for ws in wb.worksheets:
        if first_col(header(ws), DATE_COLS) is not None:
            return ws
    raise RuntimeError("no tab with an mw_date column")


def shift_bounds(day: dt.date, shift: tuple[str, str]) -> tuple[dt.datetime, dt.datetime]:
    # A shift whose start is after its end began the previous evening.
    start, end = (dt.datetime.combine(day, dt.time.fromisoformat(t)) for t in shift)
    if start >= end:
        start -= dt.timedelta(days=1)
    return start, end


def windows_in_shift(ws, start: dt.datetime, end: dt.datetime, kv_counts=None) -> list[tuple[dt.date, str, int]]:
    hdr = header(ws)
    d, t, s = first_col(hdr, DATE_COLS), first_col(hdr, TIME_COLS), first_col(hdr, STATUS_COLS)
    f = first_col(hdr, FORMAL_COLS)
    k = first_col(hdr, KV_COLS)

    def cell(row, i):
        return row[i] if i is not None and i < len(row) else None

    counts = Counter()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if f is not None:
            fm = FORMAL_RE.match(str(cell(row, f) or "").strip())
            if not fm:
                continue
            try:
                v = dt.datetime.strptime(" ".join(fm.group(1, 2, 3)), "%b %d %Y").date()
            except ValueError:  # blank dates render as "Jan 00, 1900"
                continue
            window = fm.group(4)
        else:
            v = cell(row, d)
            if isinstance(v, dt.datetime):
                v = v.date()
            window = str(cell(row, t) or "").strip()
        if not isinstance(v, dt.date) or not (start.date() <= v <= end.date()):
            continue
        status = str(cell(row, s) or "").strip().lower()
        if s is not None and (not status or status in EXCLUDED_STATUSES):
            continue
        m = re.match(r"\d{2}:\d{2}", window)
        if not m or not (start <= dt.datetime.combine(v, dt.time.fromisoformat(m.group(0))) < end):
            continue
        counts[(v, window)] += 1
        if kv_counts is not None and k is not None and str(cell(row, k) or "").strip().lower() in ("yes", "y", "true"):
            kv_counts[(v, window)] += 1
    return [(v, w, n) for (v, w), n in sorted(counts.items())]


def inspect_workbook(name: str, data: bytes, max_rows: int):
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    print(f"\n=== {name} ===")
    for ws in wb.worksheets:
        print(f"--- sheet: {ws.title} ---")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                break
            if any(c is not None for c in row):
                print(f"  r{i + 1}: {list(row)}")
    wb.close()


# ------------------------------------------------------------------- Report

def build_message(day: dt.date, sections: list[tuple[str, list[str]]]) -> str:
    out = [f"Please verify MWs for {day:%d/%m/%Y}", ""]
    for title, lines in sections:
        out += [title, SEPARATOR, *(lines or ["NA"]), ""]
    return "\n".join(out).rstrip() + "\n"


def project_sections(reader, cfg: dict, links, day: dt.date) -> list[tuple[str, list[str]]]:
    start, end = shift_bounds(day, tuple(cfg["shift_utc"]))
    projects = cfg.get("projects", {})

    def match(text):
        for i, (key, p) in enumerate(projects.items()):
            if norm(text).startswith(norm(key)):
                return i, p
        # Fall back to the closest key so a typo in the index doc ("ES Upgr4433ades")
        # still picks up that project's title and tab.
        keys = {norm(k): i for i, k in enumerate(projects)}
        close = difflib.get_close_matches(norm(text), keys, n=1, cutoff=0.8)
        if close:
            i = keys[close[0]]
            return i, list(projects.values())[i]
        return len(projects), {}

    ordered = sorted(links, key=lambda l: match(l[0])[0])
    sections = []
    for text, url in ordered:
        _, p = match(text)
        title = p.get("title") or text or url
        log(f"Reading {title} ...")
        kv_lines = []
        try:
            wb = load_workbook(io.BytesIO(reader.read(url)), read_only=True, data_only=True)
            kv = Counter()
            try:
                windows = windows_in_shift(stack_tab(wb, p.get("tab")), start, end, kv)
            finally:
                wb.close()
            lines = [
                f"{v:%b %d, %Y} {w} {n} stack{'s' if n != 1 else ''}" for v, w, n in windows
            ]
            kv_lines = [f"{v:%b %d, %Y} {w} {n:02d} stacks" for (v, w), n in sorted(kv.items())]
        except Exception as e:
            lines = [f"ERROR: could not read this sheet ({first_line(e)})"]
        sections.append((title, lines))
        if kv_lines:
            sections.append(("Kv Migration - RT", kv_lines))

    for title in cfg.get("manual_sections", []):
        sections.append((title, ["<fill in from calendar>"]))
    return sections


def copy_to_clipboard(text: str) -> bool:
    try:
        subprocess.run(["pbcopy"], input=text.encode(), check=True)
        return True
    except (OSError, subprocess.CalledProcessError):
        return False


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--date", type=dt.date.fromisoformat, default=dt.date.today(),
                    help="MW date as YYYY-MM-DD (default: today)")
    ap.add_argument("--inspect", action="store_true", help="dump each workbook's structure instead")
    ap.add_argument("--rows", type=int, default=15, help="rows per sheet with --inspect")
    args = ap.parse_args()

    cfg = json.loads(CONFIG_FILE.read_text())

    with sync_playwright() as p:
        reader = SharePointReader(p)
        try:
            log("Reading index Word doc ...")
            links = docx_hyperlinks(reader.read(cfg["index_doc"]))
            log(f"Found {len(links)} sheet link(s) in the index doc.")

            if args.inspect:
                for text, url in links:
                    try:
                        inspect_workbook(text or url, reader.read(url), args.rows)
                    except Exception as e:
                        print(f"\n=== {text or url} ===\n  FAILED: {first_line(e)}")
                return

            message = build_message(args.date, project_sections(reader, cfg, links, args.date))
        finally:
            reader.close()

    print("\n" + message)
    MESSAGES_DIR.mkdir(exist_ok=True)
    out_file = MESSAGES_DIR / f"{args.date:%Y-%m-%d}.txt"
    out_file.write_text(message)
    log(f"(Saved to {out_file})")
    if copy_to_clipboard(message):
        log("(Copied to clipboard.)")


if __name__ == "__main__":
    main()
